import argparse
import os

import openpyxl

from sffl.cli import _value_pool, main
from sffl.identity import NFL_TEAMS
from sffl.league import load_league
from sffl.render.xlsx import GROUP_GAP

DS = "sources/draftsharks.yaml"
FIXTURE = "tests/fixtures/draftsharks_value_sample.csv"


def test_render_writes_both_artifacts(tmp_path):
    pdf = str(tmp_path / "board.pdf")
    xlsx = str(tmp_path / "board.xlsx")
    rc = main(["render", "--source", DS, "--file", FIXTURE, "--year", "2026",
               "--pdf", pdf, "--xlsx", xlsx])
    assert rc == 0
    assert os.path.getsize(pdf) > 0
    assert os.path.getsize(xlsx) > 0


def test_render_requires_at_least_one_output(capsys):
    rc = main(["render", "--source", DS, "--file", FIXTURE, "--year", "2026"])
    out = capsys.readouterr().out
    assert rc == 1
    assert "--pdf" in out and "--xlsx" in out


# A dedicated fixture: draftsharks_value_sample.csv plus one extra WR row on
# team "UNS" (an unsigned free agent placeholder, not a real NFL franchise).
# The real Draft Sharks extract has 23 such rows (UNS x22, RK x1) - they have
# no bye week and build_rows correctly raises rather than rendering a blank,
# so cmd_render must filter them out itself and say how many it dropped,
# rather than silently shrinking the board or crashing.
UNROSTERED_FIXTURE = "tests/fixtures/draftsharks_render_unrostered_sample.csv"


def test_render_drops_unrostered_players_and_reports_the_count(tmp_path, capsys):
    pdf = str(tmp_path / "board.pdf")
    rc = main(["render", "--source", DS, "--file", UNROSTERED_FIXTURE, "--year", "2026",
               "--pdf", pdf])
    out = capsys.readouterr().out
    assert rc == 0
    assert "dropped 1" in out
    assert "UNS" in out


def test_dropping_unrostered_players_reprices_nobody_on_the_rendered_board(tmp_path):
    # cmd_render drops non-NFL team codes AFTER _value_pool has priced the
    # whole pool, so removing them cannot move anyone's _dollars -
    # replacement levels and the market curve are already fixed by then.
    # The ordering is deliberate, not incidental: filtering first would
    # shrink the pool the replacement level is computed from and quietly
    # reprice the entire board. This class of "obviously true" invariant
    # has broken three times in the renderers' history, so pin it.
    #
    # The comparison must cross a real boundary or it proves nothing. An
    # earlier version of this test read the dollars off the same Player
    # objects it then filtered a list of - and since filtering a Python
    # list does not mutate its elements, every assertion compared a value
    # to itself. It passed against a cmd_render deliberately mutated to
    # reprice after filtering. So: snapshot the independently-valued
    # dollars, drive the REAL `render` command end to end, and read the
    # MY$ back out of the workbook it wrote.
    args = argparse.Namespace(
        source=DS, file=UNROSTERED_FIXTURE, year=2026, set=None,
        curves=None, prices=None, policy="starter", tqb_starters=None)
    pool, _curve, _prices = _value_pool(load_league("leagues/sffl/2026.yaml"), args)

    # The workbook rounds MY$ to whole dollars, so compare on that basis.
    priced = dict((p.name, round(p.stats["_dollars"], 0)) for p in pool)
    unrostered = [p.name for p in pool if p.team not in NFL_TEAMS]
    assert unrostered, (
        "the fixture must contain at least one unrostered player, or this "
        "test filters nothing and proves nothing")

    xlsx = str(tmp_path / "board.xlsx")
    rc = main(["render", "--source", DS, "--file", UNROSTERED_FIXTURE,
               "--year", "2026", "--xlsx", xlsx])
    assert rc == 0

    ws = openpyxl.load_workbook(xlsx).active
    on_board = {}
    for col0 in (1, 1 + GROUP_GAP, 1 + 2 * GROUP_GAP):
        for r in range(1, ws.max_row + 1):
            name = ws.cell(r, col0 + 1).value
            if name in priced:
                on_board[name] = ws.cell(r, col0 + 6).value   # the MY$ column

    assert on_board, "no player from the valued pool reached the board"
    for name in unrostered:
        assert name not in on_board, (
            "%s has a non-NFL team code and must not render" % name)
    for name, my_dollars in sorted(on_board.items()):
        assert my_dollars == priced[name], (
            "%s was repriced by dropping unrostered players: the board says "
            "$%s, the valuation said $%s" % (name, my_dollars, priced[name]))
